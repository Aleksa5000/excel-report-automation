"""Turn raw sales CSV/Excel exports into a formatted, client-ready Excel report."""
import argparse
import os
import smtplib
import sys
from email.message import EmailMessage
from pathlib import Path

import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REQUIRED = ["date", "region", "product", "units", "unit_price"]
CURRENCY = '"$"#,##0.00'
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")


def load(paths) -> pd.DataFrame:
    """Read CSV/XLSX files as text so cleaning decides what is valid."""
    frames = []
    for p in map(Path, paths):
        reader = pd.read_excel if p.suffix.lower() in (".xlsx", ".xls") else pd.read_csv
        frames.append(reader(p, dtype=str))
    df = pd.concat(frames, ignore_index=True)
    df.columns = [str(c).strip().lower() for c in df.columns]
    missing = [c for c in REQUIRED if c not in df.columns]
    if missing:
        raise ValueError(f"Missing required column(s): {', '.join(missing)}")
    return df[REQUIRED]


def clean(raw: pd.DataFrame):
    """Return (valid_rows, rejected_rows, duplicates_dropped)."""
    df = raw.copy()
    for c in REQUIRED:
        df[c] = df[c].astype("string").str.strip().fillna("")
    before = len(df)
    df = df.drop_duplicates().reset_index(drop=True)
    dupes = before - len(df)

    parsed = pd.DataFrame({
        "date": pd.to_datetime(df["date"], errors="coerce"),
        "units": pd.to_numeric(df["units"], errors="coerce").astype("float64"),
        "unit_price": pd.to_numeric(df["unit_price"], errors="coerce").astype("float64"),
    })
    checks = [
        ("missing region", df["region"] == ""),
        ("missing product", df["product"] == ""),
        ("invalid date", parsed["date"].isna()),
        ("invalid units (must be a positive whole number)",
         ~(parsed["units"] > 0) | (parsed["units"] % 1 != 0)),
        ("invalid unit_price (must be > 0)", ~(parsed["unit_price"] > 0)),
    ]
    reasons = pd.Series([
        "; ".join(name for name, bad in checks if bad.iloc[i]) for i in range(len(df))
    ], index=df.index)

    ok = reasons == ""
    good = df[ok].copy()
    good["date"] = parsed.loc[ok, "date"]
    good["units"] = parsed.loc[ok, "units"].astype(int)
    good["unit_price"] = parsed.loc[ok, "unit_price"]
    good["revenue"] = (good["units"] * good["unit_price"]).round(2)
    rejected = df[~ok].assign(reason=reasons[~ok])
    return good.reset_index(drop=True), rejected.reset_index(drop=True), dupes


def kpis(good: pd.DataFrame) -> dict:
    if good.empty:
        return {}
    return {
        "Total revenue": good["revenue"].sum(),
        "Total units sold": int(good["units"].sum()),
        "Orders": len(good),
        "Average order value": good["revenue"].mean(),
        "Top product (by revenue)": good.groupby("product")["revenue"].sum().idxmax(),
        "Top region (by revenue)": good.groupby("region")["revenue"].sum().idxmax(),
        "Period": f"{good['date'].min():%Y-%m-%d} to {good['date'].max():%Y-%m-%d}",
    }


def monthly_by_region(good: pd.DataFrame) -> pd.DataFrame:
    month = good["date"].dt.strftime("%Y-%m").rename("month")
    pivot = good.pivot_table(index=month, columns="region", values="revenue", aggfunc="sum", fill_value=0)
    pivot["Total"] = pivot.sum(axis=1)
    pivot = pivot.round(2)
    return pivot.reset_index()


def top_products(good: pd.DataFrame, n: int = 10) -> pd.DataFrame:
    t = good.groupby("product").agg(units=("units", "sum"), revenue=("revenue", "sum"))
    t = t.sort_values("revenue", ascending=False).head(n).reset_index()
    t.insert(0, "rank", range(1, len(t) + 1))
    return t


def _style_table(ws, money_cols=(), filter_=True):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    for col in ws.iter_cols(min_row=2):
        if ws.cell(1, col[0].column).value in money_cols:
            for c in col:
                c.number_format = CURRENCY
    for i, col in enumerate(ws.columns, 1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col) + 3
        ws.column_dimensions[get_column_letter(i)].width = min(max(width, 12), 60)
    ws.freeze_panes = "A2"
    if filter_:
        ws.auto_filter.ref = ws.dimensions


def build_workbook(good, rejected, dupes, out_path):
    k = kpis(good)
    monthly = monthly_by_region(good) if k else pd.DataFrame({"month": []})
    top = top_products(good) if k else pd.DataFrame({"product": []})
    summary = pd.DataFrame(list(k.items()) + [
        ("Rejected rows", len(rejected)), ("Duplicates removed", dupes)], columns=["Metric", "Value"])

    with pd.ExcelWriter(out_path, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Summary", index=False)
        monthly.to_excel(xw, sheet_name="Monthly revenue by region", index=False)
        top.to_excel(xw, sheet_name="Top 10 products", index=False)
        rejected.to_excel(xw, sheet_name="Rejected rows", index=False)
        wb = xw.book

        ws = wb["Summary"]
        _style_table(ws, filter_=False)  # a filter on a KPI list is just noise
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 26
        for label, cell in ws.iter_rows(min_row=2, max_col=2):
            label.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="right")
            if label.value in ("Total revenue", "Average order value"):
                cell.number_format = CURRENCY
            elif isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"

        ws = wb["Monthly revenue by region"]
        _style_table(ws, money_cols=set(monthly.columns[1:]))
        if k:
            n_regions = len(monthly.columns) - 2  # exclude month + Total
            chart = BarChart()
            chart.type, chart.grouping, chart.overlap = "col", "stacked", 100
            chart.title = "Monthly revenue by region"
            chart.y_axis.title, chart.x_axis.title = "Revenue ($)", "Month"
            chart.y_axis.number_format = '"$"#,##0'
            chart.add_data(Reference(ws, min_col=2, max_col=1 + n_regions, min_row=1, max_row=len(monthly) + 1),
                           titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(monthly) + 1))
            chart.x_axis.delete = chart.y_axis.delete = False
            chart.width, chart.height = 26, 12
            ws.add_chart(chart, f"A{len(monthly) + 4}")

        _style_table(wb["Top 10 products"], money_cols={"revenue"})
        ws = wb["Rejected rows"]
        _style_table(ws)
        ws.column_dimensions["F"].width = 50


def send_email(path: Path) -> bool:
    """Send the report via SMTP; returns False (and does nothing) if env vars are not set."""
    env = {k: os.environ.get(k) for k in ("SMTP_HOST", "SMTP_PORT", "SMTP_USER", "SMTP_PASSWORD", "REPORT_TO")}
    if not all(env.values()):
        return False
    msg = EmailMessage()
    msg["Subject"], msg["From"], msg["To"] = f"Sales report: {path.name}", env["SMTP_USER"], env["REPORT_TO"]
    msg.set_content("Hello,\n\nPlease find the latest sales report attached.\n")
    msg.add_attachment(path.read_bytes(), maintype="application",
                       subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=path.name)
    with smtplib.SMTP(env["SMTP_HOST"], int(env["SMTP_PORT"])) as s:
        s.starttls()
        s.login(env["SMTP_USER"], env["SMTP_PASSWORD"])
        s.send_message(msg)
    return True


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Build a formatted Excel sales report from CSV/XLSX exports.")
    ap.add_argument("inputs", nargs="+", help="one or more .csv / .xlsx files")
    ap.add_argument("-o", "--output", default="report.xlsx")
    ap.add_argument("--email", action="store_true", help="email the report (needs SMTP_* and REPORT_TO env vars)")
    args = ap.parse_args(argv)

    try:
        raw = load(args.inputs)
    except (OSError, ValueError) as e:
        print(f"Error: {e}", file=sys.stderr)
        return 1
    good, rejected, dupes = clean(raw)
    build_workbook(good, rejected, dupes, args.output)

    print(f"Read {len(raw)} rows from {len(args.inputs)} file(s)")
    print(f"  {dupes} duplicates removed, {len(rejected)} rejected, {len(good)} valid")
    for name, v in kpis(good).items():
        print(f"  {name + ':':27}{v:,.2f}" if isinstance(v, float) else f"  {name + ':':27}{v}")
    print(f"Report written to {args.output}")

    if args.email:
        try:
            print("Report emailed." if send_email(Path(args.output)) else "Email skipped: SMTP env vars not set.")
        except (OSError, smtplib.SMTPException) as e:
            print(f"Email failed: {e}", file=sys.stderr)
            return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
