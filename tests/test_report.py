import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import report  # noqa: E402

ROWS = [
    ["2025-01-05", "North", "Widget", "2", "10.00"],
    ["2025-01-05", "North", "Widget", "2", "10.00"],        # exact duplicate
    ["2025-01-20", " South ", "Gadget ", "1", "30.50"],      # needs trimming
    ["2025-02-03", "North", "Gadget", "3", "30.50"],
    ["not-a-date", "East", "Widget", "1", "10"],             # bad date
    ["2025-02-04", "East", "Widget", "-1", "10"],            # bad units
    ["2025-02-05", "", "Widget", "1", "10"],                 # missing region
    ["2025-02-06", "East", "Widget", "1", "0"],              # bad price
]


@pytest.fixture
def cleaned():
    return report.clean(pd.DataFrame(ROWS, columns=report.REQUIRED))


def test_cleaning(cleaned):
    good, rejected, dupes = cleaned
    assert dupes == 1
    assert len(good) == 3
    assert set(good["region"]) == {"North", "South"}  # trimmed
    assert len(rejected) == 4
    reasons = " ".join(rejected["reason"])
    for expected in ("invalid date", "invalid units", "missing region", "invalid unit_price"):
        assert expected in reasons


def test_kpis(cleaned):
    k = report.kpis(cleaned[0])
    assert k["Total revenue"] == pytest.approx(20 + 30.5 + 91.5)
    assert k["Total units sold"] == 6
    assert k["Average order value"] == pytest.approx(142 / 3)
    assert k["Top product (by revenue)"] == "Gadget"
    assert k["Top region (by revenue)"] == "North"


def test_missing_columns(tmp_path):
    p = tmp_path / "bad.csv"
    p.write_text("date,region\n2025-01-01,North\n")
    with pytest.raises(ValueError, match="Missing required"):
        report.load([p])


def test_cli_builds_workbook(tmp_path):
    csv = tmp_path / "in.csv"
    pd.DataFrame(ROWS, columns=report.REQUIRED).to_csv(csv, index=False)
    out = tmp_path / "out.xlsx"
    assert report.main([str(csv), "-o", str(out)]) == 0
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "Monthly revenue by region", "Top 10 products", "Rejected rows"]
    assert len(wb["Monthly revenue by region"]._charts) == 1
    assert wb["Rejected rows"].max_row == 5  # header + 4
    assert wb["Top 10 products"].freeze_panes == "A2"
    assert wb["Top 10 products"].auto_filter.ref


def test_excel_input(tmp_path):
    x = tmp_path / "in.xlsx"
    pd.DataFrame(ROWS, columns=report.REQUIRED).to_excel(x, index=False)
    assert len(report.load([x])) == len(ROWS)


def test_email_skipped_without_env(monkeypatch, tmp_path):
    for k in ("SMTP_HOST", "SMTP_PORT", "SMTP_USER", "SMTP_PASSWORD", "REPORT_TO"):
        monkeypatch.delenv(k, raising=False)
    assert report.send_email(tmp_path / "x.xlsx") is False
